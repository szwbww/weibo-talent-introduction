"""Read one approved demand row and build an Apollo People Search configuration."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from core import Stop, clean


FIELDS = (
    "企业名称", "企业研究方向", "当前亟需解决的难点", "岗位名称", "岗位职责",
    "专业领域", "工作经验", "工作履历", "年龄", "薪资待遇（年薪）", "备注",
)


def normalized(value):
    return re.sub(r"\s+", "", clean(value))


def load_requirement(path, row_number, sheet=None):
    path = Path(path).expanduser().resolve()
    if not path.is_file():
        raise Stop("Excel 文件不存在。")
    if type(row_number) is not int or row_number < 1:
        raise Stop("--row 必须是有效行号。")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise Stop("Excel 文件无法读取。") from exc
    try:
        if sheet and sheet not in workbook.sheetnames:
            raise Stop("指定工作表不存在。")
        worksheet = workbook[sheet] if sheet else workbook.active
        header_row = None
        columns = {}
        for index, values in enumerate(worksheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
            current = {normalized(value): pos + 1 for pos, value in enumerate(values) if normalized(value)}
            if all(name in current for name in ("企业名称", "岗位名称", "专业领域")):
                header_row, columns = index, current
                break
        if header_row is None:
            raise Stop("前20行未找到企业名称、岗位名称、专业领域表头。")
        if row_number <= header_row:
            raise Stop("--row 必须指向表头之后的数据行。")
        requirement = {}
        for field in FIELDS:
            column = columns.get(field)
            requirement[field] = clean(worksheet.cell(row=row_number, column=column).value) if column else ""
        for field in ("企业名称", "岗位名称", "专业领域"):
            if not requirement[field]:
                raise Stop(f"Excel 第 {row_number} 行缺少 {field}。")
        requirement["source"] = {"file": str(path), "sheet": worksheet.title,
                                 "row": row_number, "header_row": header_row}
        return requirement
    finally:
        workbook.close()


def build_apollo_config(requirement, benchmark, domain):
    benchmark, domain = clean(benchmark), clean(domain).lower()
    if not benchmark or not domain:
        raise Stop("必须提供对标企业名称和域名。")
    topics = [value for value in (
        requirement.get("企业研究方向"), requirement.get("当前亟需解决的难点"),
        requirement.get("专业领域"), requirement.get("岗位职责"),
    ) if clean(value) and clean(value) != "/"]
    return {
        "strategy": "APOLLO_PEOPLE",
        "demand_company": clean(requirement["企业名称"]),
        "company": benchmark,
        "domain": domain,
        "aliases": [benchmark],
        "topics": topics,
        "titles": [clean(requirement["岗位名称"])],
        "countries": [],
        "include_historical": True,
        "overseas_only": False,
        "source_urls": [],
        "allowed_hosts": [],
        "requirement": requirement,
        "model": "deepseek-v4-flash",
        "budgets": {
            "deepseek_profile": 1, "apollo_search_pages": 1, "deepseek_rank": 1,
            "apollo_enrich": 1, "email_verify": 1, "fetch": 0, "deepseek": 0,
            "input_chars": 14000, "output_tokens": 2200,
        },
    }
